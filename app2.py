

import os
from werkzeug.utils import secure_filename
from flask import Flask, request, render_template,url_for, redirect, Blueprint,send_file
from fpdf import FPDF 
from docx import Document
from flask import send_from_directory
 
dir_path = os.path.dirname(os.path.realpath(__file__))
UPLOAD_FOLDER = 'uploads'


def convdocx():
 doc = Document()


 with open("uploads/tes_op.txt", 'r', encoding='utf-8') as file:
    doc.add_paragraph(file.read())
    doc.save("uploads/Demo1.docx")
 return("Demo1.docx")

def convpdf():
# save FPDF() class into 
# a variable pdf 
 pdf = FPDF() 

# Add a page 
 pdf.add_page() 

 pdf.add_font('gargi', '', 'gargi.ttf', uni=True) 
 pdf.set_font('gargi', '', 14)


# open the text file in read mode 
 f = open("uploads/tes_op.txt", "r", errors='replace', encoding='utf-8') 

# insert the texts in pdf 
 for x in f: 
    if x=='\n':
        x=''
    else:
        pdf.cell(200, 10, txt = x, ln = 1) 

# save the pdf with name .pdf 
 pdf.output("uploads/mypdf.pdf") 
 return("mypdf.pdf")


app_file2 = Flask(__name__) 
app_file2 = Blueprint('app_file2',__name__)
@app_file2.route("/convert")
def convert():
    return render_template('convert.html')

@app_file2.route("/fileconvert",methods=['GET','POST'])
def return_files_tut1():
    try:   
        if request.method == "POST":
         req = request.form
         userText= req.get("format")
        #userText = request.form['pdf']
        #userText1 = request.form['docx']
         print(str(userText))
         #print(str(userText1))
         if(userText=="docx"):
            res=convdocx()
            return redirect('/convertfile/'+ res)
            return render_template('convert.html')
         elif(userText=="pdf"):
            res=convpdf()
            return redirect('/convertfile/'+ res)
            return render_template('convert.html')
         else:
            return render_template('convert.html')
        return redirect(request.url)
    except Exception as e:
        return str(e)
    return None    




@app_file2.route("/convertfile/<res>", methods = ['GET'])
def convert_file(res):
    return render_template('downloads1.html',value=res)

@app_file2.route('/return-convfiles/<res>')
def return_convfiles(res):
    return send_file('uploads/'+res, as_attachment=True, attachment_filename='')
    return render_template('home.html')
    